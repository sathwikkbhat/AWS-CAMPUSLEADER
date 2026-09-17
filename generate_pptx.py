import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Initialize Presentation with 16:9 widescreen
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

blank_slide_layout = prs.slide_layouts[6]

# Sophisticated Corporate Palette (Subtle & Professional)
COLOR_DARK_BG = RGBColor(18, 25, 35)        # Deep Slate / Squid Ink
COLOR_CARD_BG = RGBColor(26, 36, 48)        # Neutral Dark Card
COLOR_CARD_BORDER = RGBColor(46, 62, 82)    # Subtle Clean Border
COLOR_AWS_ORANGE = RGBColor(236, 114, 17)   # AWS Amber / Orange
COLOR_WHITE = RGBColor(248, 250, 252)       # Clean White
COLOR_MUTED = RGBColor(148, 163, 184)       # Muted Slate Body
COLOR_SUBTLE_PILL = RGBColor(32, 44, 60)    # Subtle Pill Fill

FONT_HEADING = "Amazon Ember"
FONT_BODY = "Amazon Ember"

def set_slide_background_image(slide, img_path):
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, Inches(0), Inches(0), width=Inches(13.333), height=Inches(7.5))

def add_header(slide, title_text, category_text="AWS BUILDER CENTER"):
    # Category Pill
    tb_cat = slide.shapes.add_textbox(Inches(0.9), Inches(0.45), Inches(8), Inches(0.35))
    tf_cat = tb_cat.text_frame
    tf_cat.word_wrap = True
    tf_cat.margin_left = tf_cat.margin_right = tf_cat.margin_top = tf_cat.margin_bottom = 0
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.name = FONT_HEADING
    p_cat.font.size = Pt(10.5)
    p_cat.font.bold = True
    p_cat.font.color.rgb = COLOR_AWS_ORANGE

    # Main Title
    tb_title = slide.shapes.add_textbox(Inches(0.9), Inches(0.8), Inches(10), Inches(0.65))
    tf_title = tb_title.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_WHITE

    # AWS Logo top right
    aws_logo = "-GEO- AWS_logo_White.png"
    if os.path.exists(aws_logo):
        slide.shapes.add_picture(aws_logo, Inches(11.6), Inches(0.45), width=Inches(0.95))

# ==========================================
# SLIDE 1: TITLE SLIDE
# ==========================================
slide1 = prs.slides.add_slide(blank_slide_layout)
set_slide_background_image(slide1, "WhatsApp Image 2026-07-16 at 18.09.51.jpeg")

overlay1 = slide1.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.7), Inches(11.733), Inches(6.1)
)
overlay1.fill.solid()
overlay1.fill.fore_color.rgb = COLOR_DARK_BG
overlay1.line.color.rgb = COLOR_CARD_BORDER
overlay1.line.width = Pt(1)

if os.path.exists("-GEO- AWS_logo_White.png"):
    slide1.shapes.add_picture("-GEO- AWS_logo_White.png", Inches(10.8), Inches(1.2), width=Inches(1.2))

if os.path.exists("ABC White (1).png"):
    slide1.shapes.add_picture("ABC White (1).png", Inches(1.5), Inches(1.4), width=Inches(3.6))

# Pill Badge
badge = slide1.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(2.25), Inches(3.2), Inches(0.35)
)
badge.fill.solid()
badge.fill.fore_color.rgb = COLOR_SUBTLE_PILL
badge.line.color.rgb = COLOR_CARD_BORDER
badge.line.width = Pt(1)
tf_b = badge.text_frame
tf_b.paragraphs[0].text = "CAMPUS ANNOUNCEMENT 2026"
tf_b.paragraphs[0].font.name = FONT_HEADING
tf_b.paragraphs[0].font.size = Pt(10)
tf_b.paragraphs[0].font.bold = True
tf_b.paragraphs[0].font.color.rgb = COLOR_MUTED
tf_b.paragraphs[0].alignment = PP_ALIGN.CENTER

# Headline
tb = slide1.shapes.add_textbox(Inches(1.5), Inches(2.8), Inches(9.5), Inches(1.9))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Build What's Next with AWS"
p.font.name = FONT_HEADING
p.font.size = Pt(42)
p.font.bold = True
p.font.color.rgb = COLOR_WHITE

p2 = tf.add_paragraph()
p2.text = "Discover AWS Builder Center — Your Global Gateway to Cloud, AI & Community"
p2.font.name = FONT_BODY
p2.font.size = Pt(17)
p2.font.color.rgb = COLOR_AWS_ORANGE
p2.space_before = Pt(8)

# Presenter Profile Bar
pres_box = slide1.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(5.1), Inches(10.3), Inches(1.15)
)
pres_box.fill.solid()
pres_box.fill.fore_color.rgb = COLOR_CARD_BG
pres_box.line.color.rgb = COLOR_CARD_BORDER
pres_box.line.width = Pt(1)

tf_pres = pres_box.text_frame
p_pres1 = tf_pres.paragraphs[0]
p_pres1.text = "Presented by Sathwik Bhat"
p_pres1.font.name = FONT_HEADING
p_pres1.font.size = Pt(16)
p_pres1.font.bold = True
p_pres1.font.color.rgb = COLOR_WHITE

p_pres2 = tf_pres.add_paragraph()
p_pres2.text = "AWS Student Builder Campus Leader (SBCL)  •  @bhatsathwik  •  builder.aws.com"
p_pres2.font.name = FONT_BODY
p_pres2.font.size = Pt(12)
p_pres2.font.color.rgb = COLOR_MUTED
p_pres2.space_before = Pt(3)


# ==========================================
# SLIDE 2: MEET YOUR CAMPUS LEADER
# ==========================================
slide2 = prs.slides.add_slide(blank_slide_layout)
set_slide_background_image(slide2, "WhatsApp Image 2026-07-16 at 18.09.51.jpeg")

bg_card2 = slide2.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(11.733), Inches(6.7)
)
bg_card2.fill.solid()
bg_card2.fill.fore_color.rgb = COLOR_DARK_BG
bg_card2.line.color.rgb = COLOR_CARD_BORDER
bg_card2.line.width = Pt(1)

add_header(slide2, "Meet Your Campus Leader: Sathwik Bhat", "AWS STUDENT BUILDER CAMPUS LEADER")

# Left Column: Profile Card
left_card = slide2.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.7), Inches(4.0), Inches(5.0)
)
left_card.fill.solid()
left_card.fill.fore_color.rgb = COLOR_CARD_BG
left_card.line.color.rgb = COLOR_CARD_BORDER
left_card.line.width = Pt(1)

if os.path.exists("sathwik_circle.png"):
    slide2.shapes.add_picture("sathwik_circle.png", Inches(2.35), Inches(1.95), width=Inches(1.3), height=Inches(1.3))

tf_lc = left_card.text_frame
tf_lc.word_wrap = True
tf_lc.margin_left = tf_lc.margin_right = Inches(0.3)
tf_lc.margin_top = Inches(1.7)

p = tf_lc.paragraphs[0]
p.text = "Sathwik Bhat"
p.font.name = FONT_HEADING
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = COLOR_WHITE

p = tf_lc.add_paragraph()
p.text = "AWS Student Builder Campus Leader (SBCL)"
p.font.name = FONT_BODY
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_AWS_ORANGE
p.space_before = Pt(4)

p = tf_lc.add_paragraph()
p.text = "AWS Student Builder Campus Leader (SBCL) on our campus, helping students explore real-world cloud computing, GenAI tools, and certification guidance."
p.font.name = FONT_BODY
p.font.size = Pt(11.5)
p.font.color.rgb = COLOR_MUTED
p.space_before = Pt(14)

p = tf_lc.add_paragraph()
p.text = "Builder Center Handle:\n@bhatsathwik  •  SATHWIK K BHAT"
p.font.name = FONT_BODY
p.font.size = Pt(11.5)
p.font.bold = True
p.font.color.rgb = COLOR_WHITE
p.space_before = Pt(16)

# Right Column: 3 Pillars with clean, subtle styling
pillars = [
    ("Connecting Our Campus to AWS", "Direct bridge to AWS learning resources, developer community programs, hackathons, and technical challenges."),
    ("Hands-On Cloud & GenAI Skill Pathways", "Empowering peers across all branches and years with guided learning, next-gen AI tools, and hands-on workshops."),
    ("Mentorship & Certification Preparation", "Guiding you through AWS certification pathways (Cloud Practitioner, Solutions Architect), workshops, and project showcases.")
]

for idx, (title, desc) in enumerate(pillars):
    top_pos = Inches(1.7 + idx * 1.68)
    card = slide2.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.3), top_pos, Inches(7.0), Inches(1.48)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD_BG
    card.line.color.rgb = COLOR_CARD_BORDER
    card.line.width = Pt(1)

    tf_c = card.text_frame
    tf_c.word_wrap = True
    tf_c.margin_left = tf_c.margin_right = Inches(0.35)
    tf_c.margin_top = Inches(0.2)

    p1 = tf_c.paragraphs[0]
    p1.text = f"0{idx+1}.  {title}"
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE

    p2 = tf_c.add_paragraph()
    p2.text = desc
    p2.font.name = FONT_BODY
    p2.font.size = Pt(11)
    p2.font.color.rgb = COLOR_MUTED
    p2.space_before = Pt(4)


# ==========================================
# SLIDE 3: WHAT IS AWS BUILDER CENTER?
# ==========================================
slide3 = prs.slides.add_slide(blank_slide_layout)
set_slide_background_image(slide3, "WhatsApp Image 2026-07-16 at 18.09.50.jpeg")

bg_card3 = slide3.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(11.733), Inches(6.7)
)
bg_card3.fill.solid()
bg_card3.fill.fore_color.rgb = COLOR_DARK_BG
bg_card3.line.color.rgb = COLOR_CARD_BORDER
bg_card3.line.width = Pt(1)

add_header(slide3, "What is AWS Builder Center?", "THE BUILDER COMMUNITY PLATFORM")

tb3 = slide3.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(5.2), Inches(5.0))
tf3 = tb3.text_frame
tf3.word_wrap = True
tf3.margin_left = tf3.margin_right = tf3.margin_top = tf3.margin_bottom = 0

p = tf3.paragraphs[0]
p.text = "The Home for Builders Worldwide"
p.font.name = FONT_HEADING
p.font.size = Pt(19)
p.font.bold = True
p.font.color.rgb = COLOR_WHITE

p = tf3.add_paragraph()
p.text = "AWS Builder Center is Amazon's premier interactive space where developers, students, and cloud builders come together to learn, build, and connect."
p.font.name = FONT_BODY
p.font.size = Pt(12)
p.font.color.rgb = COLOR_MUTED
p.space_before = Pt(8)

p = tf3.add_paragraph()
p.text = "Key Highlights for Students:"
p.font.name = FONT_HEADING
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = COLOR_AWS_ORANGE
p.space_before = Pt(14)

features = [
    ("Curated Technical Articles", "Access real-world technical articles, tutorials, architecture patterns, and publish your own student projects."),
    ("Global Developer Network", "Directly interact with AWS Community Builders, AWS Heroes, and student leaders worldwide."),
    ("Hands-On Sandboxes & GenAI Tools", "Build live cloud projects, experiment with PartyRock & Amazon Bedrock, and explore serverless architectures with guided sandboxes.")
]
for f_title, f_desc in features:
    p = tf3.add_paragraph()
    p.text = f"•  {f_title}: "
    p.font.name = FONT_BODY
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.space_before = Pt(8)

    run = p.add_run()
    run.text = f_desc
    run.font.bold = False
    run.font.color.rgb = COLOR_MUTED

if os.path.exists("IMG-20260916-WA0040.jpg.jpeg"):
    img_card = slide3.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.5), Inches(1.7), Inches(5.8), Inches(5.0)
    )
    img_card.fill.solid()
    img_card.fill.fore_color.rgb = COLOR_CARD_BG
    img_card.line.color.rgb = COLOR_CARD_BORDER
    img_card.line.width = Pt(1)

    slide3.shapes.add_picture("IMG-20260916-WA0040.jpg.jpeg", Inches(6.6), Inches(2.1), width=Inches(5.6))
    
    tb_cap = slide3.shapes.add_textbox(Inches(6.6), Inches(1.75), Inches(5.5), Inches(0.3))
    tf_cap = tb_cap.text_frame
    p_cap = tf_cap.paragraphs[0]
    p_cap.text = "LIVE DASHBOARD: AWS BUILDER CENTER"
    p_cap.font.name = FONT_HEADING
    p_cap.font.size = Pt(10)
    p_cap.font.bold = True
    p_cap.font.color.rgb = COLOR_MUTED


# ==========================================
# SLIDE 4: WHY STUDENTS SHOULD JOIN (BENEFITS)
# ==========================================
slide4 = prs.slides.add_slide(blank_slide_layout)
set_slide_background_image(slide4, "WhatsApp Image 2026-07-16 at 18.09.51.jpeg")

bg_card4 = slide4.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(11.733), Inches(6.7)
)
bg_card4.fill.solid()
bg_card4.fill.fore_color.rgb = COLOR_DARK_BG
bg_card4.line.color.rgb = COLOR_CARD_BORDER
bg_card4.line.width = Pt(1)

add_header(slide4, "Why Should You Join? Student Benefits", "EXCLUSIVE PERKS FOR OUR CAMPUS")

benefits = [
    ("Curated Skill Builder Tracks", "Comprehensive on-demand curriculum and guided paths for Cloud Practitioner, Solutions Architect, and Generative AI."),
    ("Student Rewards & Exam Vouchers", "Verify your student status with SheerID to earn milestone badges, AWS cloud credits, and foundational certification exam vouchers."),
    ("Generative AI Innovation Playground", "Build working AI apps in minutes using PartyRock (Amazon Bedrock) and accelerate coding with Amazon Q Developer."),
    ("Verified Builder Profile & Network", "Showcase your real builds on a permanent builder profile (builder.aws.com/@alias) and connect with peers & mentors.")
]

for idx, (b_title, b_desc) in enumerate(benefits):
    col = idx % 2
    row = idx // 2
    x = Inches(1.0 + col * 5.7)
    y = Inches(1.7 + row * 2.1)

    card = slide4.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.4), Inches(1.9)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD_BG
    card.line.color.rgb = COLOR_CARD_BORDER
    card.line.width = Pt(1)

    tf_b = card.text_frame
    tf_b.word_wrap = True
    tf_b.margin_left = tf_b.margin_right = Inches(0.35)
    tf_b.margin_top = Inches(0.2)

    p1 = tf_b.paragraphs[0]
    p1.text = b_title
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(16)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE

    p2 = tf_b.add_paragraph()
    p2.text = b_desc
    p2.font.name = FONT_BODY
    p2.font.size = Pt(11)
    p2.font.color.rgb = COLOR_MUTED
    p2.space_before = Pt(6)

bottom_banner = slide4.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(6.1), Inches(11.1), Inches(0.7)
)
bottom_banner.fill.solid()
bottom_banner.fill.fore_color.rgb = COLOR_SUBTLE_PILL
bottom_banner.line.color.rgb = COLOR_CARD_BORDER
bottom_banner.line.width = Pt(1)

tf_bb = bottom_banner.text_frame
p_bb = tf_bb.paragraphs[0]
p_bb.text = "FRICTIONLESS STUDENT ONBOARDING  —  SIGN UP IN 60 SECONDS WITH YOUR AWS BUILDER ID"
p_bb.font.name = FONT_HEADING
p_bb.font.size = Pt(12)
p_bb.font.bold = True
p_bb.font.color.rgb = COLOR_WHITE
p_bb.alignment = PP_ALIGN.CENTER


# ==========================================
# SLIDE 5: 4-STEP SIGNUP & CLAIM ALIAS (ALL 4 SCREENSHOTS)
# ==========================================
slide5 = prs.slides.add_slide(blank_slide_layout)
set_slide_background_image(slide5, "WhatsApp Image 2026-07-16 at 18.09.50.jpeg")

bg_card5 = slide5.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(11.733), Inches(6.7)
)
bg_card5.fill.solid()
bg_card5.fill.fore_color.rgb = COLOR_DARK_BG
bg_card5.line.color.rgb = COLOR_CARD_BORDER
bg_card5.line.width = Pt(1)

add_header(slide5, "How to Sign Up in 4 Quick Steps (Takes 60s)", "QUICK ONBOARDING WALKTHROUGH")

# 4 Steps with real screenshots
steps = [
    ("Step 01", "Scan & Tap Sign In", "Open link & tap 'Sign in' on top", "step1_cropped.jpg"),
    ("Step 02", "Tap < Back Arrow", "Click Back arrow to access handle setup", "step2_cropped.jpg"),
    ("Step 03", "Claim Unique @Alias", "Choose handle & select Country: India", "Screenshot_20260916-191419.png"),
    ("Step 04", "Confirm Student Status", "Select 'Yes', degree & graduation year", "Screenshot_20260916-191446.png"),
]

for idx, (s_num, s_title, s_desc, img_file) in enumerate(steps):
    x = Inches(1.0 + idx * 2.85)
    
    step_card = slide5.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.65), Inches(2.65), Inches(5.15)
    )
    step_card.fill.solid()
    step_card.fill.fore_color.rgb = COLOR_CARD_BG
    step_card.line.color.rgb = COLOR_CARD_BORDER
    step_card.line.width = Pt(1)

    tb_s = slide5.shapes.add_textbox(x, Inches(1.75), Inches(2.65), Inches(0.95))
    tf_s = tb_s.text_frame
    tf_s.word_wrap = True
    tf_s.margin_left = tf_s.margin_right = Inches(0.15)
    
    p = tf_s.paragraphs[0]
    p.text = s_num.upper()
    p.font.name = FONT_HEADING
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_AWS_ORANGE
    p.alignment = PP_ALIGN.CENTER

    p = tf_s.add_paragraph()
    p.text = s_title
    p.font.name = FONT_HEADING
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(2)

    img_y = Inches(2.8)
    if img_file and os.path.exists(img_file):
        slide5.shapes.add_picture(img_file, x + Inches(0.32), img_y, width=Inches(2.01), height=Inches(3.85))


# ==========================================
# SLIDE 6: CALL TO ACTION (QR CODE & LINK)
# ==========================================
slide6 = prs.slides.add_slide(blank_slide_layout)
set_slide_background_image(slide6, "WhatsApp Image 2026-07-16 at 18.09.51.jpeg")

bg_card6 = slide6.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(11.733), Inches(6.7)
)
bg_card6.fill.solid()
bg_card6.fill.fore_color.rgb = COLOR_DARK_BG
bg_card6.line.color.rgb = COLOR_CARD_BORDER
bg_card6.line.width = Pt(1)

add_header(slide6, "Scan Now & Join AWS Builder Center", "JOIN OUR CAMPUS COMMUNITY")

# Left Box: QR Code Display Card (Clean White)
qr_card = slide6.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.7), Inches(4.5), Inches(5.0)
)
qr_card.fill.solid()
qr_card.fill.fore_color.rgb = RGBColor(255, 255, 255)
qr_card.line.color.rgb = RGBColor(230, 235, 242)
qr_card.line.width = Pt(1)

if os.path.exists("image.png"):
    slide6.shapes.add_picture("image.png", Inches(1.4), Inches(1.95), width=Inches(3.7))

tb_qr_lbl = slide6.shapes.add_textbox(Inches(1.0), Inches(5.85), Inches(4.5), Inches(0.6))
tf_qr_lbl = tb_qr_lbl.text_frame
p_ql = tf_qr_lbl.paragraphs[0]
p_ql.text = "SCAN WITH YOUR PHONE CAMERA"
p_ql.font.name = FONT_HEADING
p_ql.font.size = Pt(11)
p_ql.font.bold = True
p_ql.font.color.rgb = RGBColor(26, 36, 48)
p_ql.alignment = PP_ALIGN.CENTER

# Right Column: Clean CTA, Unique Link, Sathwik's Info
right_card = slide6.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.8), Inches(1.7), Inches(6.5), Inches(5.0)
)
right_card.fill.solid()
right_card.fill.fore_color.rgb = COLOR_CARD_BG
right_card.line.color.rgb = COLOR_CARD_BORDER
right_card.line.width = Pt(1)

tf_rc = right_card.text_frame
tf_rc.word_wrap = True
tf_rc.margin_left = tf_rc.margin_right = Inches(0.4)
tf_rc.margin_top = Inches(0.35)

p = tf_rc.paragraphs[0]
p.text = "Exclusive Sign-Up Link:"
p.font.name = FONT_HEADING
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = COLOR_MUTED

p = tf_rc.add_paragraph()
p.text = "https://bit.ly/4hQ40lk"
p.font.name = FONT_HEADING
p.font.size = Pt(22)
p.font.bold = True
p.font.color.rgb = COLOR_AWS_ORANGE
p.space_before = Pt(4)

p = tf_rc.add_paragraph()
p.text = "Why you need to claim it today:"
p.font.name = FONT_HEADING
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = COLOR_WHITE
p.space_before = Pt(16)

cta_points = [
    ("First-come, first-served handles", "Unique @aliases cannot be changed once chosen. Reserve your personal name or handle today."),
    ("Instant access to student perks", "Unlock comprehensive Skill Builder curriculum, GenAI tools & milestone badges right away."),
    ("Connect with Sathwik Bhat (@bhatsathwik)", "Connect with me on campus for project guidance, student study groups, and certifications.")
]

for ct_title, ct_desc in cta_points:
    p = tf_rc.add_paragraph()
    p.text = f"•  {ct_title}: "
    p.font.name = FONT_BODY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.space_before = Pt(7)
    run = p.add_run()
    run.text = ct_desc
    run.font.bold = False
    run.font.color.rgb = COLOR_MUTED

p = tf_rc.add_paragraph()
p.text = "AWS Student Builder Campus Leader (SBCL)"
p.font.name = FONT_BODY
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_AWS_ORANGE
p.space_before = Pt(18)

output_path = "AWS_Builder_Center_Class_Announcement.pptx"
prs.save(output_path)
print(f"Successfully generated: {output_path}")
